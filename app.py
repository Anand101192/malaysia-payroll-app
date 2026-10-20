import streamlit as st
import pandas as pd
import openpyxl
import io

st.set_page_config(page_title="Malaysia Automated Payroll System", page_icon="🇲🇾", layout="wide")

st.title("🇲🇾 Malaysia Web Payroll Engine")
st.markdown("""
**Drop & Calculate System**: Drop your monthly attendance & overtime Excel sheet. 
The engine extracts all staff members, merges OT hours, displays all employee and employer deduction/contribution columns, formats days cleanly without decimals, and generates an **auto-calculating Excel report**.
""")

# --- FILE UPLOADER ---
uploaded_file = st.file_uploader("Drop Attendance Excel File Here", type=["xlsx", "xls"])

if uploaded_file is not None:
    try:
        xls = pd.ExcelFile(uploaded_file)
        
        all_sheets = xls.sheet_names
        filtered_sheets = [s for s in all_sheets if 'AGENCY' not in s.upper() and 'SHEET1' not in s.upper()]
        if not filtered_sheets:
            filtered_sheets = all_sheets
            
        default_index = 0
        for i, s in enumerate(filtered_sheets):
            if 'ATTENDANCE (2)' in s.upper():
                default_index = i
                break
                
        selected_sheet = st.sidebar.selectbox("Select Attendance Sheet", filtered_sheets, index=default_index)
        
        # Read selected sheet with header row 8
        df = pd.read_excel(uploaded_file, sheet_name=selected_sheet, header=8)
        df.columns = [str(c).strip() for c in df.columns]
        
        name_col = 'STAFF / CREW NAME' if 'STAFF / CREW NAME' in df.columns else next((c for c in df.columns if 'NAME' in c.upper() and c != 'RESIGN/TERMINATE'), df.columns[2])
        pos_col = 'POSITION' if 'POSITION' in df.columns else None
        ic_col = 'STAFF / CREW' if 'STAFF / CREW' in df.columns else None
        total_col = 'TOTAL' if 'TOTAL' in df.columns else None
        
        df_clean = df.dropna(subset=[name_col]).copy()
        
        # READ OVERTIME SHEET FOR OT HOURS
        ot_data = {}
        if 'OVERTIME' in xls.sheet_names:
            try:
                ot_df = pd.read_excel(uploaded_file, sheet_name='OVERTIME', header=5)
                ot_df.columns = [str(c).strip() for c in ot_df.columns]
                ot_name_col = 'NAME' if 'NAME' in ot_df.columns else ot_df.columns[2]
                
                for _, ot_row in ot_df.iterrows():
                    ot_name = str(ot_row.get(ot_name_col, '')).strip().upper()
                    if ot_name and ot_name != 'NAN' and 'NAME' not in ot_name:
                        try:
                            val1 = ot_row.iloc[-3]
                            val2 = ot_row.iloc[-2]
                            val3 = ot_row.iloc[-1]
                            
                            h1 = float(val1) if pd.notna(val1) else 0.0
                            h2 = float(val2) if pd.notna(val2) else 0.0
                            h3 = float(val3) if pd.notna(val3) else 0.0
                            ot_data[ot_name] = {'ot_1.5': h1, 'ot_2.0': h2, 'ot_3.0': h3}
                        except:
                            pass
            except Exception as e:
                pass

        default_base = st.sidebar.number_input("Default Base Salary for Crew (RM)", value=2200.0, step=100.0)
        manager_base = st.sidebar.number_input("Default Base Salary for Manager (RM)", value=4500.0, step=100.0)
        
        records = []
        for idx, row in df_clean.iterrows():
            name = str(row.get(name_col, '')).strip()
            
            if not name or name.lower() == 'nan' or 'NAME' in name.upper() or 'NO.' in name.upper() or name.isdigit() or len(name) < 2:
                continue
                
            position = str(row.get(pos_col, 'Staff')).strip() if pos_col and pos_col in df.columns else 'Staff'
            ic_no = str(row.get(ic_col, '-')).strip() if ic_col and ic_col in df.columns else '-'
            
            raw_days = row.get(total_col, 26) if total_col and total_col in df.columns else 26
            try:
                days_worked = float(raw_days)
            except:
                days_worked = 26.0
                
            is_manager = "MANAGER" in position.upper()
            base = manager_base if is_manager else default_base
            
            ot_info = ot_data.get(name.upper(), {'ot_1.5': 0.0, 'ot_2.0': 0.0, 'ot_3.0': 0.0})
            
            # Python calculations
            earned_basic_val = round((base / 26.0) * min(days_worked, 26.0), 2)
            hourly_rate = base / 26.0 / 8.0
            ot_pay_val = round((ot_info['ot_1.5'] * hourly_rate * 1.5) + 
                               (ot_info['ot_2.0'] * hourly_rate * 2.0) + 
                               (ot_info['ot_3.0'] * hourly_rate * 3.0), 2)
            
            # Employee Deductions
            epf_val = round(earned_basic_val * 0.11, 2)
            socso_val = round(min(earned_basic_val, 5000.0) * 0.005, 2)
            socso_24h_val = round(earned_basic_val * 0.0075, 2)
            eis_val = round(min(earned_basic_val, 5000.0) * 0.002, 2)
            pcb_val = round((earned_basic_val - 3100.0) * 0.03, 2) if earned_basic_val > 3100.0 else 0.0
            total_ded_val = round(epf_val + socso_val + socso_24h_val + eis_val + pcb_val, 2)
            basic_after_ded_val = round(earned_basic_val - total_ded_val, 2)
            net_pay_val = round(basic_after_ded_val + ot_pay_val, 2)
            
            # Employer Contributions
            epf_er_val = round(earned_basic_val * (0.13 if earned_basic_val <= 5000 else 0.12), 2)
            socso_er_val = round(min(earned_basic_val, 5000.0) * 0.0175, 2)
            eis_er_val = round(min(earned_basic_val, 5000.0) * 0.002, 2)
            
            records.append({
                'Staff Name': name,
                'IC / ID': ic_no if ic_no.lower() != 'nan' else '-',
                'Position': position if position.lower() != 'nan' else 'Staff',
                'Days Worked': int(round(days_worked)),  # Clean whole number
                'Base Salary (RM)': f"{base:.2f}",
                'Earned Basic (RM)': f"{earned_basic_val:.2f}",
                'EPF Employee (RM)': f"{epf_val:.2f}",
                'SOCSO Employee (RM)': f"{socso_val:.2f}",
                'SOCSO 24 Jam perlindungan (RM)': f"{socso_24h_val:.2f}",
                'EIS Employee (RM)': f"{eis_val:.2f}",
                'PCB Tax (RM)': f"{pcb_val:.2f}",
                'Total Deductions (RM)': f"{total_ded_val:.2f}",
                'Basic Salary After Deductions (RM)': f"{basic_after_ded_val:.2f}",
                'OT Hours (1.5x)': int(round(ot_info['ot_1.5'])),
                'OT Hours (2.0x)': int(round(ot_info['ot_2.0'])),
                'OT Hours (3.0x)': int(round(ot_info['ot_3.0'])),
                'Total OT Pay (RM)': f"{ot_pay_val:.2f}",
                'Net Salary Payable (RM)': f"{net_pay_val:.2f}",
                'EPF Employer (RM)': f"{epf_er_val:.2f}",
                'SOCSO Employer (RM)': f"{socso_er_val:.2f}",
                'EIS Employer (RM)': f"{eis_er_val:.2f}"
            })
            
        display_df = pd.DataFrame(records)
        
        if len(display_df) == 0:
            st.warning("No staff rows extracted. Please verify sheet selection.")
        else:
            st.success(f"Success! Extracted **{len(display_df)}** staff records with employer sections and clean day numbering from `{selected_sheet}`.")
            
            st.subheader("📊 Summary Payroll Breakdown")
            col1, col2, col3, col4 = st.columns(4)
            col1.metric("Total Staff Count", f"{len(display_df)}")
            col2.metric("Total Basic Earned", f"RM {display_df['Earned Basic (RM)'].astype(float).sum():,.2f}")
            col3.metric("Total OT Payout", f"RM {display_df['Total OT Pay (RM)'].astype(float).sum():,.2f}")
            col4.metric("Total Net Payable", f"RM {display_df['Net Salary Payable (RM)'].astype(float).sum():,.2f}")
            
            st.subheader("📋 Calculated Payroll Data")
            st.dataframe(display_df, use_container_width=True)
            
            # Build Excel workbook with live formulas and proper cell formatting
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Payroll Summary"
            
            headers = list(display_df.columns)
            ws.append(headers)
            
            for i, r in display_df.iterrows():
                row_num = i + 2
                row_formulas = [
                    r['Staff Name'],
                    r['IC / ID'],
                    r['Position'],
                    int(r['Days Worked']),
                    float(r['Base Salary (RM)']),
                    f'=IF(D{row_num}>26, (E{row_num}/26)*26, (E{row_num}/26)*D{row_num})',
                    f'=F{row_num}*0.11',
                    f'=MIN(F{row_num}, 5000)*0.005',
                    f'=F{row_num}*0.0075',
                    f'=MIN(F{row_num}, 5000)*0.002',
                    f'=IF(F{row_num}>3100, (F{row_num}-3100)*0.03, 0)',
                    f'=SUM(G{row_num}:K{row_num})',
                    f'=F{row_num}-L{row_num}',
                    int(r['OT Hours (1.5x)']),
                    int(r['OT Hours (2.0x)']),
                    int(r['OT Hours (3.0x)']),
                    f'=(N{row_num}*(E{row_num}/26/8)*1.5) + (O{row_num}*(E{row_num}/26/8)*2.0) + (P{row_num}*(E{row_num}/26/8)*3.0)',
                    f'=M{row_num}+Q{row_num}',
                    f'=IF(F{row_num}<=5000, F{row_num}*0.13, F{row_num}*0.12)',
                    f'=MIN(F{row_num}, 5000)*0.0175',
                    f'=MIN(F{row_num}, 5000)*0.002'
                ]
                ws.append(row_formulas)
                
            # Apply number formats in Excel
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=4, max_col=ws.max_column):
                for col_idx, cell in enumerate(row, start=4):
                    if col_idx in [4, 14, 15, 16]:  # Days worked and OT hours (Whole numbers)
                        cell.number_format = '#,##0'
                    else:  # Monetary columns
                        cell.number_format = '0.00'

            output = io.BytesIO()
            wb.save(output)
            processed_data = output.getvalue()
            
            st.download_button(
                label="📥 Download Auto-Calculating Editable Payroll Report (.xlsx)",
                data=processed_data,
                file_name=f"Editable_Payroll_{selected_sheet}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        
    except Exception as e:
        st.error(f"Error processing file: {e}")
